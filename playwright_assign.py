"""WhatsApp Message Sender + Smart Data Extractor.

Assignment file: playwright_assign.py
Use this only to contact people who have agreed to receive your messages.
"""

from __future__ import annotations

import json
import random
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from playwright.sync_api import (
        BrowserContext,
        Locator,
        Page,
        Playwright,
        TimeoutError as PlaywrightTimeoutError,
        sync_playwright,
    )
except ImportError as exc:
    print("[ERROR] A required package is missing.")
    print("Run: pip install playwright openpyxl")
    print("Then run: playwright install chromium")
    raise SystemExit(1) from exc


# ---------------------------------------------------------------------------
# Configuration - adjust these values only when necessary.
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
SCREENSHOTS_DIR = OUTPUT_DIR / "screenshots"
PROFILE_DIR = BASE_DIR / "whatsapp_profile"
WHATSAPP_URL = "https://web.whatsapp.com/"

HEADLESS = False  # WhatsApp QR login requires a visible browser.
ACTION_TIMEOUT_MS = 30_000
LOGIN_TIMEOUT_MS = 300_000
MIN_DELAY_MS = 2_000
MAX_DELAY_MS = 5_000
EXTRACT_MESSAGE_COUNT = 3

# WhatsApp Web changes occasionally. Each list contains safe selector fallbacks.
SEARCH_SELECTORS = [
    'input[placeholder*="Search or start a new chat" i]',
    'input[placeholder*="Search" i]',
    '[role="searchbox"]',
    '[aria-label*="Search or start a new chat" i]',
    'div[contenteditable="true"][data-tab="3"]',
    'div[contenteditable="true"][role="textbox"][aria-placeholder*="Search" i]',
    'div[contenteditable="true"][aria-label*="Search input textbox" i]',
    'div[contenteditable="true"][aria-label*="Search" i]',
    'div[contenteditable="true"][title*="Search" i]',
]

# The chat pane is a more stable login indicator than the search field alone.
LOGIN_READY_SELECTORS = [
    "#pane-side",
    '[data-testid="chat-list"]',
    'div[aria-label="Chat list"]',
    *SEARCH_SELECTORS,
]
MESSAGE_BOX_SELECTORS = [
    'textarea[placeholder*="Type a message" i]',
    'div[contenteditable="true"][aria-placeholder*="Type a message" i]',
    'div[contenteditable="true"][aria-label*="Type a message" i]',
    'footer div[contenteditable="true"][data-tab="10"]',
    'footer div[contenteditable="true"][aria-label*="message" i]',
    'footer div[contenteditable="true"][role="textbox"]',
]
CHAT_PANEL_SELECTORS = ["#main", 'main[role="main"]']


def show_status(level: str, message: str) -> None:
    """Print a clear timestamped status message."""
    print(f"[{datetime.now():%H:%M:%S}] [{level}] {message}")


def human_pause(page: Page) -> None:
    """Wait for a random 2-5 seconds as required by the assignment."""
    page.wait_for_timeout(random.randint(MIN_DELAY_MS, MAX_DELAY_MS))


def prepare_folders() -> None:
    """Create output folders safely when they are missing."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)
    PROFILE_DIR.mkdir(parents=True, exist_ok=True)


def normalize_phone(value: Any) -> str:
    """Return a clean international phone number or raise a useful error."""
    if value is None:
        raise ValueError("Phone is empty")

    # Excel may return a whole number as a float, such as 919876543210.0.
    if isinstance(value, float) and value.is_integer():
        value = int(value)

    raw_phone = str(value).strip()
    compact_phone = re.sub(r"[\s().-]", "", raw_phone)
    if not re.fullmatch(r"\+\d{7,15}", compact_phone):
        raise ValueError("Phone must contain +, country code, and 7-15 digits")
    return compact_phone


def read_contacts(path: Path) -> list[dict[str, str]]:
    """Read and validate Name, Phone, and Message columns from contacts.xlsx."""
    if not path.is_file():
        raise FileNotFoundError(
            f"Input file not found: {path}\n"
            "Create contacts.xlsx with columns: Name, Phone, Message"
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            raise ValueError("contacts.xlsx is empty")

        headers = {
            str(value).strip().lower(): index
            for index, value in enumerate(header_row)
            if value is not None
        }
        required_headers = {"name", "phone", "message"}
        missing = required_headers - set(headers)
        if missing:
            raise ValueError(f"Missing Excel columns: {', '.join(sorted(missing))}")

        contacts: list[dict[str, str]] = []
        for excel_row, row in enumerate(rows, start=2):
            name_value = row[headers["name"]] if headers["name"] < len(row) else None
            phone_value = row[headers["phone"]] if headers["phone"] < len(row) else None
            message_value = (
                row[headers["message"]] if headers["message"] < len(row) else None
            )

            name = "" if name_value is None else str(name_value).strip()
            message = "" if message_value is None else str(message_value).strip()
            try:
                phone = normalize_phone(phone_value)
                if not name:
                    raise ValueError("Name is empty")
                contacts.append(
                    {
                        "row": str(excel_row),
                        "name": name,
                        "phone": phone,
                        "message_template": message,
                        "validation_error": "",
                    }
                )
            except ValueError as exc:
                contacts.append(
                    {
                        "row": str(excel_row),
                        "name": name,
                        "phone": "" if phone_value is None else str(phone_value),
                        "message_template": message,
                        "validation_error": str(exc),
                    }
                )
    finally:
        workbook.close()

    if not contacts:
        raise ValueError("contacts.xlsx has headers but no contact rows")
    return contacts


def first_visible_locator(page: Page, selectors: list[str], timeout_ms: int) -> Locator:
    """Return the first visible element found from a list of selector fallbacks."""
    time_each = max(1_000, timeout_ms // len(selectors))
    for selector in selectors:
        try:
            # Explicit wait_for_selector is required by the assignment.
            page.wait_for_selector(selector, state="visible", timeout=time_each)
            locator = page.locator(selector).first
            if locator.is_visible():
                return locator
        except PlaywrightTimeoutError:
            continue
    raise PlaywrightTimeoutError(f"No visible element matched: {selectors}")


def wait_for_login(page: Page) -> None:
    """Open WhatsApp and continuously check for the logged-in chat interface."""
    show_status("INFO", "Opening WhatsApp Web. Scan the QR code if requested.")
    page.goto(WHATSAPP_URL, wait_until="domcontentloaded", timeout=60_000)

    deadline = time.monotonic() + (LOGIN_TIMEOUT_MS / 1_000)
    while time.monotonic() < deadline:
        # #pane-side is WhatsApp's chat-list container after a successful login.
        chat_pane = page.locator("#pane-side")
        if chat_pane.count() and chat_pane.first.is_visible():
            show_status("SUCCESS", "WhatsApp Web login and chat list detected.")
            return

        # Some versions expose the search textbox before #pane-side is detected.
        for selector in SEARCH_SELECTORS:
            candidate = page.locator(selector).first
            if candidate.count() and candidate.is_visible():
                show_status("SUCCESS", "WhatsApp Web login detected.")
                return

        page.wait_for_timeout(1_000)

    diagnostic = OUTPUT_DIR / "login_timeout.png"
    page.screenshot(path=str(diagnostic), full_page=True)
    raise RuntimeError(
        "WhatsApp login was not detected within 5 minutes. "
        f"A diagnostic screenshot was saved at: {diagnostic}"
    )


def find_search_box(page: Page, timeout_ms: int = ACTION_TIMEOUT_MS) -> Locator:
    """Find the visible chat-search textbox without relying on one data-tab value."""
    deadline = time.monotonic() + (timeout_ms / 1_000)
    search_buttons = [
        'button[aria-label*="Search" i]',
        '[role="button"][aria-label*="Search" i]',
        '[data-icon="search"]',
    ]

    while time.monotonic() < deadline:
        for selector in SEARCH_SELECTORS:
            candidate = page.locator(selector).first
            if candidate.count() and candidate.is_visible():
                return candidate

        # If WhatsApp shows a search icon first, click it to reveal the textbox.
        for selector in search_buttons:
            button = page.locator(selector).first
            if button.count() and button.is_visible():
                button.click()
                page.wait_for_timeout(500)
                break

        # Final fallback: choose a visible textbox in the upper-left chat panel.
        textboxes = page.locator(
            'input[type="text"], input[type="search"], [role="searchbox"], '
            'div[contenteditable="true"][role="textbox"]'
        )
        for index in range(textboxes.count()):
            candidate = textboxes.nth(index)
            box = candidate.bounding_box() if candidate.is_visible() else None
            if box and box["x"] < 700 and box["y"] < 250:
                return candidate
        page.wait_for_timeout(500)

    raise PlaywrightTimeoutError("The WhatsApp chat search box was not found")


def clear_and_type(page: Page, locator: Locator, text: str) -> None:
    """Clear and type without keeping a stale WhatsApp input element."""
    position = locator.bounding_box()
    if not position:
        raise PlaywrightTimeoutError("The detected input field is not visible")

    # Click the live centre of the detected element. These are measured values,
    # not fixed coordinates, so different screen sizes are supported.
    page.mouse.click(
        position["x"] + position["width"] / 2,
        position["y"] + position["height"] / 2,
    )
    page.wait_for_timeout(300)
    page.keyboard.press("Control+A")
    page.keyboard.press("Backspace")
    if text:
        page.keyboard.type(text, delay=random.randint(35, 90))


def find_message_box(page: Page, timeout_ms: int = ACTION_TIMEOUT_MS) -> Locator:
    """Find the message box after a contact conversation opens."""
    deadline = time.monotonic() + (timeout_ms / 1_000)
    while time.monotonic() < deadline:
        for selector in MESSAGE_BOX_SELECTORS:
            candidate = page.locator(selector).last
            if candidate.count() and candidate.is_visible():
                return candidate

        # Fallback for new layouts: locate an editable field in the lower-right.
        fields = page.locator(
            'textarea, input[type="text"], '
            'div[contenteditable="true"][role="textbox"], '
            'div[contenteditable="true"]'
        )
        viewport = page.viewport_size or {"width": 1440, "height": 900}
        for index in range(fields.count()):
            candidate = fields.nth(index)
            position = candidate.bounding_box() if candidate.is_visible() else None
            if not position:
                continue
            if (
                position["x"] > viewport["width"] * 0.30
                and position["y"] > viewport["height"] * 0.55
            ):
                return candidate
        page.wait_for_timeout(500)

    raise PlaywrightTimeoutError("The Type a message box was not found")


def open_contact(page: Page, name: str, phone: str) -> bool:
    """Search by phone/name, click the visible result, and open the chat."""
    for search_value in (phone, name):
        search_box = find_search_box(page)
        clear_and_type(page, search_box, search_value)
        human_pause(page)

        # Excel may contain "sarath" while WhatsApp displays "Sarath".
        # Match the complete name without considering upper/lowercase.
        exact_name = re.compile(rf"^\s*{re.escape(name)}\s*$", re.IGNORECASE)
        matching_elements = page.get_by_text(exact_name)
        visible_matches: list[tuple[float, Locator, dict[str, float]]] = []

        for index in range(matching_elements.count()):
            result = matching_elements.nth(index)
            if not result.is_visible():
                continue
            position = result.bounding_box()
            if not position:
                continue

            # The contact result must be inside the left search-results area.
            if position["x"] <= 650 and position["y"] >= 100:
                area = position["width"] * position["height"]
                visible_matches.append((area, result, position))

        # A small text element is safer than a large page/container match.
        visible_matches.sort(key=lambda item: item[0])
        for _area, result, position in visible_matches:
            show_status("INFO", f"Visible contact result found: {name}")
            page.mouse.click(
                position["x"] + position["width"] / 2,
                position["y"] + position["height"] / 2,
            )
            page.wait_for_timeout(1_000)

            try:
                find_message_box(page, 20_000)
                show_status("SUCCESS", f"Conversation opened for {name}.")
                human_pause(page)
                return True
            except PlaywrightTimeoutError:
                continue

        # Keyboard fallback selects the first visible search result.
        search_box = find_search_box(page)
        search_position = search_box.bounding_box()
        if not search_position:
            continue
        page.mouse.click(
            search_position["x"] + search_position["width"] / 2,
            search_position["y"] + search_position["height"] / 2,
        )
        page.keyboard.press("ArrowDown")
        page.wait_for_timeout(500)
        page.keyboard.press("Enter")
        try:
            find_message_box(page, 20_000)
            human_pause(page)
            return True
        except PlaywrightTimeoutError:
            search_box = find_search_box(page)
            clear_and_type(page, search_box, "")
            continue

    # Fallback for a valid WhatsApp number that is not saved as a contact.
    digits = re.sub(r"\D", "", phone)
    page.goto(
        f"{WHATSAPP_URL}send?phone={quote(digits)}",
        wait_until="domcontentloaded",
        timeout=60_000,
    )
    try:
        find_message_box(page, 30_000)
        human_pause(page)
        return True
    except PlaywrightTimeoutError:
        return False

def send_message(page: Page, message: str) -> None:
    """Type, send and confirm the outgoing WhatsApp message."""

    # Find the message box.
    message_box = find_message_box(page, ACTION_TIMEOUT_MS)

    # Count existing outgoing messages before sending.
    outgoing_messages = page.locator("#main div.message-out")
    previous_message_count = outgoing_messages.count()

    # Get the current message-box position.
    message_position = message_box.bounding_box()

    if not message_position:
        raise PlaywrightTimeoutError("The message box is not visible")

    # Click the detected message box.
    page.mouse.click(
        message_position["x"] + message_position["width"] / 2,
        message_position["y"] + message_position["height"] / 2,
    )

    # Type the personalized message.
    page.keyboard.type(
        message,
        delay=random.randint(35, 90),
    )

    # Send immediately after typing.
    page.keyboard.press("Enter")

    # Wait until WhatsApp displays a new outgoing message.
    deadline = time.monotonic() + (ACTION_TIMEOUT_MS / 1000)

    while time.monotonic() < deadline:
        if outgoing_messages.count() > previous_message_count:
            latest_message = outgoing_messages.last

            if latest_message.is_visible():
                show_status("SUCCESS", "The outgoing message was confirmed.")
                return

        page.wait_for_timeout(500)

    raise PlaywrightTimeoutError(
        "The message was typed, but a new outgoing message was not confirmed"
    )

def extract_last_incoming_messages(page: Page, limit: int = 3) -> list[str]:
    """Extract the latest visible incoming message texts from the open chat."""
    incoming = page.locator("#main div.message-in")
    count = incoming.count()
    extracted: list[str] = []

    for index in range(max(0, count - limit), count):
        bubble = incoming.nth(index)
        # copyable-text contains the displayed chat text in current WhatsApp Web.
        text_nodes = bubble.locator("span.selectable-text.copyable-text")
        parts = [text_nodes.nth(i).inner_text().strip() for i in range(text_nodes.count())]
        message_text = "\n".join(part for part in parts if part)
        if not message_text:
            message_text = bubble.inner_text().strip()
        if message_text:
            extracted.append(message_text)
    return extracted[-limit:]


def safe_filename(value: str) -> str:
    """Convert a contact name into a safe screenshot filename."""
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")
    return cleaned[:50] or "contact"


def take_sent_screenshot(page: Page, name: str, row_number: str) -> Path:
    """Save a screenshot of the open chat after the message is confirmed."""
    filename = f"row_{row_number}_{safe_filename(name)}_{datetime.now():%H%M%S}.png"
    screenshot_path = SCREENSHOTS_DIR / filename
    chat_panel = first_visible_locator(page, CHAT_PANEL_SELECTORS, ACTION_TIMEOUT_MS)
    chat_panel.screenshot(path=str(screenshot_path))
    if not screenshot_path.is_file() or screenshot_path.stat().st_size == 0:
        raise RuntimeError("Screenshot file was not created correctly")
    return screenshot_path


def take_error_screenshot(page: Page, name: str, row_number: str) -> str:
    """Save the current browser screen when one contact fails.

    The screenshot helps the user see whether WhatsApp displayed an invalid
    number, changed its page layout, or failed to load the conversation.
    """
    filename = (
        f"ERROR_row_{row_number}_{safe_filename(name)}_"
        f"{datetime.now():%H%M%S}.png"
    )
    screenshot_path = SCREENSHOTS_DIR / filename
    try:
        page.screenshot(path=str(screenshot_path), full_page=True)
        if screenshot_path.is_file() and screenshot_path.stat().st_size > 0:
            return str(screenshot_path.relative_to(BASE_DIR))
    except Exception:
        # Do not hide the original automation error if a screenshot also fails.
        pass
    return ""


def new_result(contact: dict[str, str]) -> dict[str, Any]:
    """Create a standard report entry for one contact."""
    return {
        "excel_row": int(contact["row"]),
        "name": contact["name"],
        "phone": contact["phone"],
        "personalized_message": "",
        "sent_status": "not_processed",
        "sent_at": "",
        "screenshot": "",
        "last_3_incoming_messages": [],
        "error": "",
    }


def process_contact(page: Page, contact: dict[str, str]) -> dict[str, Any]:
    """Validate, open, message, capture, and extract data for one contact."""
    result = new_result(contact)
    name = contact["name"]

    if contact["validation_error"]:
        result["sent_status"] = "invalid_input"
        result["error"] = contact["validation_error"]
        return result

    if not contact["message_template"]:
        result["sent_status"] = "skipped_blank_message"
        result["error"] = "Message cell is blank; no default text was invented"
        return result

    # Official assignment format: {name}. For beginner convenience, also
    # accept the actual name inside braces, such as {sarath} for Sarath.
    placeholder_pattern = rf"\{{(?:name|{re.escape(name)})\}}"
    message = re.sub(
        placeholder_pattern,
        lambda _match: name,
        contact["message_template"],
        flags=re.IGNORECASE,
    )
    result["personalized_message"] = message

    try:
        show_status("INFO", f"Searching for {name} ({contact['phone']})")
        if not open_contact(page, name, contact["phone"]):
            result["sent_status"] = "contact_not_found"
            result["error"] = "No matching search result was found"
            return result

        send_message(page, message)
        result["sent_status"] = "sent"
        result["sent_at"] = datetime.now().astimezone().isoformat(timespec="seconds")

        screenshot_path = take_sent_screenshot(page, name, contact["row"])
        result["screenshot"] = str(screenshot_path.relative_to(BASE_DIR))
        result["last_3_incoming_messages"] = extract_last_incoming_messages(
            page, EXTRACT_MESSAGE_COUNT
        )
        show_status("SUCCESS", f"Message sent and recorded for {name}.")
    except PlaywrightTimeoutError as exc:
        result["sent_status"] = "failed"
        result["error"] = f"A required WhatsApp element timed out: {exc}"
        result["screenshot"] = take_error_screenshot(
            page, name, contact["row"]
        )
        show_status("ERROR", f"Timed out while processing {name}.")
    except Exception as exc:  # Continue safely with the remaining contacts.
        result["sent_status"] = "failed"
        result["error"] = f"{type(exc).__name__}: {exc}"
        result["screenshot"] = take_error_screenshot(
            page, name, contact["row"]
        )
        show_status("ERROR", f"Could not process {name}: {exc}")
    return result


def save_json_report(results: list[dict[str, Any]], path: Path) -> None:
    """Save full report details as readable UTF-8 JSON."""
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_file": CONTACTS_FILE.name,
        "total_contacts": len(results),
        "results": results,
    }
    with path.open("w", encoding="utf-8") as file:
        json.dump(report, file, ensure_ascii=False, indent=2)


def save_excel_report(results: list[dict[str, Any]], path: Path) -> None:
    """Save a formatted summary report as Excel."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "WhatsApp Summary"
    headers = [
        "Excel Row", "Name", "Phone", "Personalized Message", "Sent Status",
        "Sent At", "Screenshot", "Last 3 Incoming Messages", "Error",
    ]
    sheet.append(headers)

    for result in results:
        sheet.append(
            [
                result["excel_row"], result["name"], result["phone"],
                result["personalized_message"], result["sent_status"],
                result["sent_at"], result["screenshot"],
                "\n---\n".join(result["last_3_incoming_messages"]), result["error"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="9C1C2C")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 12, "B": 22, "C": 18, "D": 42, "E": 24, "F": 28,
              "G": 42, "H": 55, "I": 48}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(path)
    workbook.close()


def validate_reports(json_path: Path, excel_path: Path, expected_rows: int) -> None:
    """Reopen both reports and confirm they contain the expected row count."""
    if not json_path.is_file() or json_path.stat().st_size == 0:
        raise RuntimeError("JSON report was not created")
    if not excel_path.is_file() or excel_path.stat().st_size == 0:
        raise RuntimeError("Excel report was not created")

    with json_path.open("r", encoding="utf-8") as file:
        json_data = json.load(file)
    if len(json_data.get("results", [])) != expected_rows:
        raise RuntimeError("JSON report row count is incorrect")

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        summary_rows = workbook["WhatsApp Summary"].max_row - 1
    finally:
        workbook.close()
    if summary_rows != expected_rows:
        raise RuntimeError("Excel report row count is incorrect")


def launch_browser(playwright: Playwright) -> BrowserContext:
    """Launch Chromium with a dedicated profile that keeps WhatsApp login."""
    return playwright.chromium.launch_persistent_context(
        user_data_dir=str(PROFILE_DIR),
        headless=HEADLESS,
        viewport={"width": 1440, "height": 900},
        locale="en-US",
    )


def confirm_responsible_use() -> None:
    """Require confirmation that all recipients consented to the messages."""
    print("\nUse this bot only for contacts who agreed to receive these messages.")
    answer = input("Have all contacts agreed to receive messages? (yes/no): ").strip().lower()
    if answer not in {"yes", "y", "consent"}:
        raise RuntimeError(
            "Confirmation was not accepted. Type yes to start the automation."
        )


def main() -> int:
    """Run the complete WhatsApp automation and reporting workflow."""
    contacts: list[dict[str, str]] = []
    results: list[dict[str, Any]] = []
    run_error = ""
    try:
        prepare_folders()
        contacts = read_contacts(CONTACTS_FILE)
        show_status("INFO", f"Loaded {len(contacts)} contact row(s).")
        confirm_responsible_use()

        with sync_playwright() as playwright:
            context = launch_browser(playwright)
            try:
                page = context.pages[0] if context.pages else context.new_page()
                page.set_default_timeout(ACTION_TIMEOUT_MS)
                wait_for_login(page)

                for contact in contacts:
                    results.append(process_contact(page, contact))
            finally:
                # Close the browser while Playwright's event loop is still active.
                try:
                    context.close()
                    show_status("INFO", "Browser closed safely.")
                except Exception as close_error:
                    show_status("WARNING", f"Browser close warning: {close_error}")
    except KeyboardInterrupt:
        show_status("WARNING", "Stopped safely by the user (Ctrl+C).")
        run_error = "Stopped safely by the user"
    except Exception as exc:
        show_status("ERROR", str(exc))
        run_error = str(exc)

    # Produce reports even when login or browser automation fails.
    if contacts:
        processed_rows = {result["excel_row"] for result in results}
        for contact in contacts:
            if int(contact["row"]) not in processed_rows:
                result = new_result(contact)
                result["sent_status"] = "failed"
                result["error"] = run_error or "Automation stopped before this row"
                results.append(result)

        try:
            report_date = date.today().isoformat()
            json_path = OUTPUT_DIR / f"whatsapp_report_{report_date}.json"
            excel_path = OUTPUT_DIR / f"whatsapp_report_{report_date}.xlsx"
            save_json_report(results, json_path)

            try:
                save_excel_report(results, excel_path)
            except PermissionError:
                # Excel locks an open workbook on Windows. Preserve the report by
                # saving a new timestamped file instead of losing the run results.
                timestamp = datetime.now().strftime("%H%M%S")
                excel_path = OUTPUT_DIR / (
                    f"whatsapp_report_{report_date}_{timestamp}.xlsx"
                )
                show_status(
                    "WARNING",
                    "The dated Excel report is open or locked. "
                    f"Saving this run as {excel_path.name} instead.",
                )
                save_excel_report(results, excel_path)

            validate_reports(json_path, excel_path, len(contacts))
            show_status("SUCCESS", f"JSON report verified: {json_path}")
            show_status("SUCCESS", f"Excel report verified: {excel_path}")
        except Exception as report_error:
            show_status("ERROR", f"Report creation failed: {report_error}")
            return 1

    return 1 if run_error else 0


if __name__ == "__main__":
    sys.exit(main())
